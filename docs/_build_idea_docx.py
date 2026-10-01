# -*- coding: utf-8 -*-
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent / "SofraGaza-Complete-Idea.docx"
OUT_AR = Path(__file__).resolve().parent / (
    "\u0633\u0641\u0631\u0629 \u063a\u0632\u0629 \u2014 \u0648\u062b\u064a\u0642\u0629 \u0627\u0644\u0641\u0643\u0631\u0629 \u0627\u0644\u0643\u0627\u0645\u0644\u0629.docx"
)

PRIMARY = RGBColor(0xB4, 0x23, 0x18)
DARK = RGBColor(0x1C, 0x19, 0x17)
MUTED = RGBColor(0x57, 0x53, 0x4E)
GOLD = RGBColor(0xB4, 0x53, 0x09)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CREAM = "F5EDE6"
ROW_ALT = "FAF7F4"


def set_run_rtl(run, font="Traditional Arabic", size=14, bold=False, color=None, italic=False):
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), font)
    rFonts.set(qn("w:hAnsi"), font)
    rFonts.set(qn("w:cs"), font)
    rFonts.set(qn("w:eastAsia"), font)
    rtl = OxmlElement("w:rtl")
    rPr.append(rtl)
    cs = OxmlElement("w:cs")
    rPr.append(cs)


def set_para_rtl(p, align="right", space_after=8, space_before=0, line=1.15):
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = line
    p.alignment = {
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    pPr.append(bidi)
    jc = pPr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        pPr.append(jc)
    jc.set(qn("w:val"), "right" if align in ("right", "justify") else align)


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "E7E5E4")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def add_text(p, text, **kwargs):
    run = p.add_run(text)
    set_run_rtl(run, **kwargs)
    return run


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    set_para_rtl(p, "right", space_after=10, space_before=18 if level == 1 else 12)
    if level == 1:
        add_text(p, text, size=22, bold=True, color=PRIMARY)
        bar = doc.add_paragraph()
        set_para_rtl(bar, "right", space_after=12, space_before=0)
        add_text(bar, "━" * 28, size=10, color=PRIMARY)
    elif level == 2:
        add_text(p, text, size=16, bold=True, color=DARK)
    else:
        add_text(p, text, size=14, bold=True, color=GOLD)
    return p


def para(doc, text, size=13.5, space=8, bold=False, color=None, align="justify"):
    p = doc.add_paragraph()
    set_para_rtl(p, align, space_after=space)
    add_text(p, text, size=size, bold=bold, color=color or DARK)
    return p


def bullets(doc, items, numbered=False):
    for i, item in enumerate(items, 1):
        p = doc.add_paragraph()
        set_para_rtl(p, "justify", space_after=4, space_before=1)
        mark = f"{i}.  " if numbered else "•  "
        add_text(p, mark, size=13, bold=True, color=PRIMARY)
        add_text(p, item, size=13, color=DARK)


def callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade_cell(cell, CREAM)
    set_cell_border(cell)
    p1 = cell.paragraphs[0]
    set_para_rtl(p1, "right", space_after=4)
    add_text(p1, title, size=13, bold=True, color=PRIMARY)
    p2 = cell.add_paragraph()
    set_para_rtl(p2, "justify", space_after=2)
    add_text(p2, body, size=12.5, color=DARK)
    spacer = doc.add_paragraph()
    set_para_rtl(spacer, "right", space_after=6)
    add_text(spacer, "", size=6)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_cell(cell, "B42318")
        set_cell_border(cell)
        p = cell.paragraphs[0]
        set_para_rtl(p, "center", space_after=2, space_before=2)
        add_text(p, h, size=12, bold=True, color=WHITE, font="Traditional Arabic")
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.rows[r + 1].cells[c]
            if r % 2 == 1:
                shade_cell(cell, ROW_ALT)
            set_cell_border(cell)
            p = cell.paragraphs[0]
            set_para_rtl(p, "right", space_after=2, space_before=2)
            add_text(p, val, size=12, color=DARK)
    spacer = doc.add_paragraph()
    set_para_rtl(spacer, "right", space_after=8)
    add_text(spacer, "", size=4)


def page_break(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run._r.append(br)


def set_doc_rtl(doc):
    styles = doc.styles["Normal"]
    styles.font.name = "Traditional Arabic"
    styles.font.size = Pt(13.5)
    rPr = styles.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), "Traditional Arabic")
    rFonts.set(qn("w:hAnsi"), "Traditional Arabic")
    rFonts.set(qn("w:cs"), "Traditional Arabic")
    sect = doc.sections[0]
    sect.page_width = Cm(21.0)
    sect.page_height = Cm(29.7)
    sect.right_margin = Cm(2.0)
    sect.left_margin = Cm(2.0)
    sect.top_margin = Cm(1.8)
    sect.bottom_margin = Cm(1.8)
    sectPr = sect._sectPr
    pgSz = sectPr.find(qn("w:pgSz"))
    if pgSz is not None:
        pgSz.set(qn("w:orient"), "portrait")
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    sectPr.append(bidi)


def footer(doc):
    section = doc.sections[0]
    fp = section.footer.paragraphs[0]
    set_para_rtl(fp, "center", space_after=0)
    add_text(fp, "سفرة غزة  ·  وثيقة الفكرة الكاملة  ·  سريّ للاستخدام الداخلي", size=9, color=MUTED)
    pPr = fp._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), "6")
    top.set(qn("w:color"), "B42318")
    pBdr.append(top)
    pPr.append(pBdr)


def build():
    doc = Document()
    set_doc_rtl(doc)
    footer(doc)

    # Cover
    for _ in range(3):
        p = doc.add_paragraph()
        set_para_rtl(p, "center", space_after=0)
        add_text(p, " ", size=14)

    k = doc.add_paragraph()
    set_para_rtl(k, "center", space_after=6)
    add_text(k, "منصة الضيافة وتوصيل الطعام الأولى في قطاع غزة", size=14, bold=True, color=GOLD)

    t = doc.add_paragraph()
    set_para_rtl(t, "center", space_after=8)
    add_text(t, "سفرة غزة", size=40, bold=True, color=PRIMARY)

    s = doc.add_paragraph()
    set_para_rtl(s, "center", space_after=18)
    add_text(s, "Sofra Gaza", size=16, color=MUTED, font="Calibri")

    line = doc.add_paragraph()
    set_para_rtl(line, "center", space_after=18)
    add_text(line, "━━━━━━━━━━━━━━━━━━━━", size=12, color=PRIMARY)

    sub = doc.add_paragraph()
    set_para_rtl(sub, "center", space_after=8)
    add_text(sub, "وثيقة الفكرة الكاملة", size=22, bold=True, color=DARK)

    desc = doc.add_paragraph()
    set_para_rtl(desc, "center", space_after=28)
    add_text(desc, "شرح مفصّل للمنتج، الأدوار، الرحلات، والخدمات — بلغة الفكرة لا بلغة البرمجة", size=13, color=MUTED)

    callout(
        doc,
        "الغرض من هذه الوثيقة",
        "تعريف سفرة غزة كما يفهمها صاحب الفكرة، الشريك، والمستثمر: ماذا نقدّم، لمن، وكيف تشتغل المنظومة من طلب الوجبة حتى تسليمها، ومن انضمام المطعم حتى ظهوره للزبائن.",
    )

    meta = doc.add_paragraph()
    set_para_rtl(meta, "center", space_after=4, space_before=24)
    add_text(meta, "غزة، فلسطين", size=12, color=MUTED)
    meta2 = doc.add_paragraph()
    set_para_rtl(meta2, "center")
    add_text(meta2, "أكتوبر 2026", size=12, color=MUTED)

    page_break(doc)

    heading(doc, "فهرس المحتويات")
    toc = [
        "1. المقدمة والرؤية",
        "2. المشكلة التي نحلّها",
        "3. الفكرة باختصار",
        "4. من هم المستخدمون",
        "5. تجربة الزبون",
        "6. رحلة الطلب من السلة حتى باب البيت",
        "7. الطلب الجماعي",
        "8. الولاء: النقاط، المستويات، والعضويات",
        "9. المحفظة والدفع المحلي",
        "10. النمو العضوي: ادعُ صديق",
        "11. تجربة صاحب المطعم أو الكافي",
        "12. الإعلان المدفوع للمطاعم",
        "13. تجربة مندوب التوصيل (الكابتن)",
        "14. لوحة الإدارة وفريق العمل",
        "15. المناطق ورسوم التوصيل",
        "16. النموذج المالي",
        "17. ما الذي يظهر وما الذي يتوقف",
        "18. المبادئ التي لا نتنازل عنها",
        "19. الخلاصة",
    ]
    bullets(doc, toc, numbered=False)

    page_break(doc)

    heading(doc, "1. المقدمة والرؤية")
    para(
        doc,
        "سفرة غزة منصة ضيافة وتوصيل طعام مبنية لأهل غزة أولاً: مطاعم وكافيهات محلية، طرق دفع يفهمها السوق، ومندوبون يوصلون الوجبة وهي سخنة. الاسم نفسه يقول الفكرة: السفرة مكان يجتمع عليه الناس، والمنصة تجمع المطعم والزبون والكابتن في حركة واحدة.",
    )
    para(
        doc,
        "الرؤية ليست «تطبيق توصيل عام». الرؤية منصة وطنية تدعم المطابخ الغزية، تسهّل الطلب على العائلة والأصدقاء، وتُبقي المال يدور داخل الاقتصاد المحلي قدر الإمكان: جوال باي، بال باي، وتحويل بنك فلسطين، بدل الاعتماد على بوابات لا تناسب الواقع.",
    )
    bullets(
        doc,
        [
            "للزبون: طلب سهل، مناطق توصيل واضحة، نقاط ومكافآت، وطلب جماعي مع الأصحاب.",
            "للمطعم: ظهور بعد اعتماد الإدارة، إدارة منيو وطلبات، وخيار إعلان مدفوع للصدارة.",
            "للكابتن: عمل منظم، مستحقات شفافة، وصرف عبر الإدارة.",
            "للإدارة: تشغيل يومي كامل من الاعتماد حتى التسوية المالية.",
        ],
    )

    heading(doc, "2. المشكلة التي نحلّها")
    para(doc, "الطلب من مطاعم غزة اليوم يتوزّع على واتساب، اتصالات، ومعرفة شخصية. هذا يُتعب المطعم، ويضيّع على الزبون وقت المقارنة، وعلى المجموعة حساب «مين دفع كم».")
    bullets(
        doc,
        [
            "الزبون يدور من الصفر كل مرة، بدون ذاكرة لمفضّلاته وطلباته السابقة.",
            "الأصدقاء يطلبون من نفس المطعم بطرق متفرّقة: فواتير كثيرة وتوصيل مكرر.",
            "المطعم يحتاج اعتماد وثقة أمام الجمهور، ومكان واحد يستقبل الطلبات.",
            "الدفع يحتاج قنوات محلية يثق فيها الناس، مع مراجعة إشعار التحويل.",
            "النمو الحقيقي يأتي من صديق يدعو صديقاً، لا من حملة إطلاق لمرة واحدة.",
        ],
    )

    heading(doc, "3. الفكرة باختصار")
    para(
        doc,
        "سفرة غزة سوق إلكتروني للطعام في قطاع غزة. الزبون يتصفّح المطاعم حسب المنطقة والصنف، يضيف للسلة، يدفع من المحفظة أو بتحويل محلي مع صورة الإشعار، ويتابع طلبه حتى التسليم. المطعم لا يظهر للناس إلا بعد موافقة الإدارة، وبعد الموافقة يبقى ظاهراً حتى توقفُه الإدارة بنفسها — ليس باشتراك ظهور ينتهي لوحده.",
    )
    callout(
        doc,
        "جملة المنتج",
        "شهيتك المفضلة، تصلك بأقصى سرعة — من مطبخ غزّي معتمد إلى عتبة دارك، بنقاط ومكافآت وطرق دفع يفهمها البلد.",
    )

    heading(doc, "4. من هم المستخدمون")
    para(doc, "المنصة أربعة عوالم واضحة. كل عالم له مدخل، وصلاحيات، وما يقدر يشتريه أو يديره. التسوق والسلة والعضوية مخصصة للزبون فقط؛ صاحب المطعم والإدارة لا يدخلون كزبائن من نفس الحساب.")
    add_table(
        doc,
        ["الدور", "ماذا يفعل", "كيف يدخل"],
        [
            ["الزبون", "يتصفّح، يطلب، يدفع، يجمع نقاط، يدعو أصدقاء", "حساب بجوال فلسطيني بعد رمز تحقق"],
            ["صاحب المطعم / الكافي", "يدير المنيو والطلبات ويمكنه شراء إعلان", "طلب انضمام ثم اعتماد الإدارة"],
            ["مندوب التوصيل (الكابتن)", "يستلم الطلبات المعيَّنة ويسلّمها ويطلب صرف أرباحه", "تسجيل ثم موافقة الإدارة"],
            ["الإدارة", "تشغّل المنصة: اعتماد، طلبات، مالية، فريق", "حساب مدير؛ الأعلى يضيف طاقماً بصلاحيات"],
        ],
    )
    para(
        doc,
        "رقم الجوال هو الهوية. تسجيل الزبون يبدأ بالتحقق برسالة، ثم الاسم وكلمة المرور. أرقام جوّال وأوريدو ضمن النمط المحلي مقبولة.",
        size=13,
    )

    heading(doc, "5. تجربة الزبون")
    heading(doc, "5.1 الاكتشاف", 2)
    para(
        doc,
        "الصفحة الرئيسية تعرض التصنيفات (شاورما، بيتزا، مشاوي، حلويات، بحريات، كافيهات…) والمطاعم الجاهزة للتوصيل الآن حسب المنطقة المختارة. البحث يقترح مطاعم وأطباقاً أثناء الكتابة. الزبون يغيّر حي التوصيل في أي وقت، ورسوم التوصيل تظهر وفق الحي.",
    )
    heading(doc, "5.2 صفحة المطعم والسلة", 2)
    para(
        doc,
        "داخل المطعم تظهر الأصناف المتاحة مع السعر وإمكانية التخصيص عند الحاجة. القلب يحفظ المطعم أو الطبق في المفضلة للوصول السريع لاحقاً. السلة لمطعم واحد في الطلب العادي، مع كوبون خصم إن وُجد.",
    )
    heading(doc, "5.3 الحساب الشخصي", 2)
    bullets(
        doc,
        [
            "طلباتي: متابعة الحالة والفاتورة وإلغاء ما يزال قابلاً للإلغاء.",
            "المفضلة: مطاعم وأصناف محفوظة.",
            "عناويني: مواقع التوصيل المتكررة.",
            "النقاط: الرصيد وسجل الحركات واستبدال نقاط بوجبات أو مشروبات.",
            "المحفظة: الرصيد، الشحن بتحويل محلي، والسجل.",
            "ادعُ صديق: كود الدعوة الدائم.",
            "الإشعارات: تأكيد الطلب، النقاط، الدعوة، وحالة التحويل.",
        ],
    )
    heading(doc, "5.4 الذاكرة والاقتراح", 2)
    para(
        doc,
        "بدل ما يبدأ الزبون من الصفر كل زيارة، تظهر على الرئيسية فقرة «بناءً على طلباتك السابقة، جرب هذا»: أطباق من مطاعم ما زالت ظاهرة، مستندة لطلباته المحتسبة. المفضلة تكمّل الصورة: قلب على المطعم أو الطبق يكفي لحفظه.",
    )

    heading(doc, "6. رحلة الطلب من السلة حتى باب البيت")
    para(doc, "الطلب يمرّ على لوحة واضحة يراها المطعم والإدارة، ويمكن سحب البطاقة من عمود لعمود لتحديث الحالة.")
    add_table(
        doc,
        ["المرحلة", "المعنى"],
        [
            ["بانتظار التأكيد", "الطلب وصل وينتظر قبول المطعم أو الإدارة"],
            ["مؤكد", "قُبل الطلب ويبدأ التحضير"],
            ["قيد التحضير", "المطبخ يجهّز الطلب"],
            ["قيد التوصيل", "الكابتن في الطريق"],
            ["تم التسليم", "وصل الزبون؛ تُحتسب النقاط ومستحقات الكابتن"],
            ["مرفوض / ملغي", "لم يكتمل الطلب؛ يُعاد ما يلزم للمحفظة حسب الحالة"],
        ],
    )
    para(
        doc,
        "بعد التسليم يمكن للزبون تقييم المطعم. التقييمات تمر على الإدارة قبل الظهور العام حتى تبقى الجودة تحت السيطرة.",
    )

    heading(doc, "7. الطلب الجماعي")
    para(
        doc,
        "أكثر من شخص يريدون نفس المطعم، في نفس الوقت، بدون فوضى. واحد يفتح «طلب جماعي»، يضيف أرقام أصحاب مسجَّلين كزبائن، وكل واحد يختار أصنافه ويدفع نصيبه. النتيجة: فاتورة واحدة للمطعم، توصيل واحد لعنوان المضيف، وتوزيع عادل للدفع.",
    )
    bullets(
        doc,
        [
            "المضيف يختار المطعم ويدعو الأرقام.",
            "كل عضو يدخل حصته من المنيو، أو يعتذر وينسحب.",
            "نافذة التجميع محدودة بساعات حتى لا يبقى الطلب معلّقاً.",
            "عند الجاهزية يدفع كل مشارك حصته (محفظة أو تحويل).",
            "من اعتذر لا يُحسب عليه شيء، ولا يدخل في التوصيل.",
            "عنوان التسليم عنوان المضيف؛ رحلة الكابتن واحدة.",
        ],
    )
    callout(
        doc,
        "لماذا هذا مهم في غزة؟",
        "السفرة عندنا جماعية أصلاً: مكتب، بيت عائلة، مجموعة أصحاب. الطلب الجماعي يقلّل التوصيل المكرر والتكلفة، ويخلي الحساب واضحاً بين الناس.",
    )

    heading(doc, "8. الولاء: النقاط، المستويات، والعضويات")
    heading(doc, "8.1 النقاط", 2)
    para(
        doc,
        "مع كل طلب مكتمل يكسب الزبون نقاطاً حسب قيمة الطعام (ويمكن احتساب التوصيل حسب إعداد الإدارة). الإدارة تضبط «كل كم شيكل = نقطة». المطعم يمكن أن يكون له سعر نقاط خاص، والصنف يمكن أن يكون له عدد نقاط ثابت. النقاط تُستبدل بمشروب أو وجبة من المنيو حسب التسعير المعتمد.",
    )
    heading(doc, "8.2 مستويات الإنفاق", 2)
    para(doc, "كل ما صرف الزبون أكثر على طلبات محتسبة، يرتقي مستوى حسابه تلقائياً:")
    add_table(
        doc,
        ["المستوى", "عتبة الإنفاق تقريباً", "المزية الرئيسية"],
        [
            ["برونزي", "من 150 ₪", "بداية المسار ونقاط أوضح"],
            ["فضي", "من 300 ₪", "نقاط بمعدّل أعلى قليلاً وأولوية خدمة"],
            ["ذهبي", "من 600 ₪", "نقاط 1.5× وأولوية تحضير وتوصيل"],
            ["بلاتيني (VIP)", "من 1,000 ₪", "نقاط مضاعفة 2× وأعلى أولوية"],
        ],
    )
    heading(doc, "8.3 العضويات الرقمية", 2)
    para(
        doc,
        "العضوية اختيار مدفوع شهري، بسعرها الحقيقي المعروض، وبعد تحويل ومراجعة الإدارة تُفعَّل. ليست تخفيضاً وهمياً على السعر.",
    )
    add_table(
        doc,
        ["العضوية", "السعر الشهري", "ماذا تعطي"],
        [
            ["الأساسية", "50 ₪", "خصم 5٪ على الطلب، ونقاط بمعدّل 1.25"],
            ["المميزة", "100 ₪", "خصم 10٪، توصيل مجاني، ونقاط بمعدّل 1.5"],
        ],
    )
    para(
        doc,
        "يمكن طلب بطاقة عضوية رقمية/مادية حسب مسار الإدارة. قرب انتهاء الاشتراك يصل تنبيه للزبون حتى يجدّد ولا يفقد الخصم.",
    )

    heading(doc, "9. المحفظة والدفع المحلي")
    para(
        doc,
        "الدفع مصمَّم لواقع غزة: إما رصيد محفظة داخل سفرة، أو تحويل مباشر مع رفع صورة الإشعار. الإدارة تراجع الحوالة وتوافق أو ترفض مع سبب واضح.",
    )
    bullets(
        doc,
        [
            "محفظة جوال باي: رقم واسم مستفيد وباركود تضبطهم الإدارة.",
            "محفظة بال باي (PalPay): نفس الفكرة لقناة ثانية يثق فيها الزبون.",
            "بنك فلسطين: اسم البنك، رقم الحساب، الآيبان، واسم المستفيد.",
            "ملاحظة توجيهية تظهر في شاشات الدفع: كتابة رقم الهاتف أو الطلب في بيان التحويل.",
        ],
    )
    para(
        doc,
        "شحن المحفظة نفس القنوات. بعد الاعتماد يصير الرصيد جاهزاً للطلب التالي بدون انتظار إشعار جديد. الطلب من المحفظة يُخصم فوراً؛ وإذا أُلغي أو رُفض يُعاد ما يلزم وفق سياسة التشغيل.",
    )

    heading(doc, "10. النمو العضوي: ادعُ صديق")
    para(
        doc,
        "هذا برنامج دائم، ليس حملة إطلاق تُغلق بعدها. كل زبون يملك كود دعوة فريداً وصفحة «ادعُ صديق، وكلاكما ياخذ نقاط». يرسل الكود أو رابط التسجيل. الصديق الجديد، عند إنشاء حسابه، يجد خيار «شخص دعاك؟ حط كوده».",
    )
    bullets(
        doc,
        [
            "إذا صحّ الكود: الطرفان يأخذان نقاطاً فوراً (الافتراضي 50 و50، قابلاً للتعديل من الإدارة).",
            "الكود يُستخدم مرة لكل حساب جديد؛ لا يمكن للشخص استخدام كود نفسه.",
            "كل صديق جديد إضافي يعطي الداعي نقاطاً من جديد — هذا محرّك النمو.",
            "كود صاحب مطعم أو إدارة لا يُحتسب؛ الدعوة بين الزبائن.",
            "رابط التسجيل يمكن أن يحمل الكود مسبقاً حتى لا يضيّعه الصديق أثناء التحقق من الجوال.",
        ],
    )

    heading(doc, "11. تجربة صاحب المطعم أو الكافي")
    para(
        doc,
        "صاحب المكان يسجّل طلب انضمام: بيانات المكان، المنطقة، الرخصة إن وُجدت، الهوية، وساعات العمل. الحالة الأولى «جاري التحقق». خلالها يقدر يجهّز المنيو والصور، لكن المكان لا يظهر في الرئيسية ولا يُطلب منه.",
    )
    bullets(
        doc,
        [
            "الإدارة توافق فينشر فوراً للزبائن، أو ترفض مع سبب يمكن تصحيحه وإعادة الإرسال.",
            "بعد الموافقة يبقى ظاهراً باستمرار، بلا عدّاد أيام يطفئه لوحده.",
            "الإدارة وحدها توقف اللوحة وتخفي المكان إن لزم؛ البيانات تبقى محفوظة لحين إعادة الفتح.",
            "من اللوحة: المنيو، حالة الطلبات الحية، الإشعارات، وطلب إعلان مدفوع.",
            "التسوق كزبون ليس من حساب المطعم؛ الحساب للشراكة والتشغيل.",
        ],
    )

    heading(doc, "12. الإعلان المدفوع للمطاعم")
    para(
        doc,
        "الظهور العادي بعد الاعتماد مجاني ومستمر. الصدارة اختيار تجاري: المطعم يحدد عدد الأيام، يُحسب السعر (20 شيكلاً لليوم)، يحوّل ويرفع الإشعار. بعد موافقة الإدارة يظهر في المقدمة طوال المدة المدفوعة، ثم يعود للترتيب الطبيعي. الإعلان لا يشتري «حق البقاء على المنصة»؛ البقاء أصلاً مضمون بعد الاعتماد.",
    )

    heading(doc, "13. تجربة مندوب التوصيل (الكابتن)")
    para(
        doc,
        "الكابتن يسجّل بياناته ونوع المركبة وصوراً للتحقق. بعد موافقة الإدارة يستلم الطلبات التي تُعيَّن له، ويتابعها: استلام من المطعم، توصيل، تسليم. أرباحه مربوطة برسوم التوصيل للطلبات المسلَّمة: المنصة تحتفظ بنسبة صغيرة من رسم التوصيل، والكابتن يستلم الباقي. يطلب صرف المستحقات، والإدارة توافق أو ترفض مع متابعة الرصيد المتاح.",
    )
    bullets(
        doc,
        [
            "لوحة للكابتن: الطلبات النشطة، المنتهية، والمحفظة/الأرباح.",
            "تعيين الطلب من الإدارة؛ يمكن فك التعيين عند الحاجة.",
            "الشفافية أهم من التعقيد: الكابتن يعرف ماذا يستحق قبل ما يطلب الصرف.",
        ],
    )

    heading(doc, "14. لوحة الإدارة وفريق العمل")
    para(
        doc,
        "هذه غرفة عمليات المنصة. المدير الأعلى يملك اللوحة كاملة، ويضيف مدراء فرعيين كلٌّ حسب وحدات محددة: طلبات، توصيل، مطاعم، مالية، زبائن، عضويات، تقييمات، كوبونات، محتوى الرئيسية، إعدادات، وسجل تعديلات.",
    )
    heading(doc, "14.1 التشغيل اليومي", 2)
    bullets(
        doc,
        [
            "اعتماد أو رفض مطاعم ومندوبين.",
            "لوحة طلبات حية بالسحب والإفلات بين الأعمدة.",
            "تعيين كابتن للطلب ومتابعة التوصيل.",
            "مراجعة حوالات المحفظة والعضويات والإعلانات.",
            "كوبونات خصم، تقييمات، وبنرات الرئيسية وشركاء الواجهة.",
            "بحث شامل في الطلبات والمطاعم والزبائن.",
        ],
    )
    heading(doc, "14.2 سجل التعديلات", 2)
    para(
        doc,
        "كل إجراء مهم يُسجَّل باسم من فعله وتاريخه ووقته: مثلاً «محمد المزيني — اعتماد مطعم الزيتون — في يوم وساعة محددة». هذا ليس للتجميل؛ هو ثقة داخل الفريق ومساءلة عند الخلاف. المحاولات الفاشلة (مثل رقم مكرر) لا تُحسب تعديلاً حتى لا يمتلئ السجل بالضجيج.",
    )

    heading(doc, "15. المناطق ورسوم التوصيل")
    para(
        doc,
        "التوصيل مربوط بأحياء حقيقية، والإدارة تقدر تضيف منطقة أو تعدّل الرسم بدون إعادة بناء المنصة. الرسم الافتراضي موجود، ولكل حي سعره إن حُدّد.",
    )
    add_table(
        doc,
        ["المنطقة", "الرسم الاسترشادي"],
        [
            ["الرمال، تل الهوى، النصر، البلدة القديمة، الميناء", "10 ₪"],
            ["الشجاعية", "12 ₪"],
            ["دير البلح", "15 ₪"],
            ["خانيونس", "20 ₪"],
        ],
    )
    para(doc, "العضوية المميزة تجعل التوصيل مجانياً للزبون المشترك طالما اشتراكه ساري.")

    heading(doc, "16. النموذج المالي")
    para(
        doc,
        "المال يدخل من تشغيل حقيقي لا من إجبار المطعم على تجديد ظهوره. بعد اعتماد المطعم يبقى في السوق؛ الإيراد يأتي من العمولة والرسوم والخدمات الاختيارية والعضويات.",
    )
    add_table(
        doc,
        ["المصدر", "الفكرة"],
        [
            ["عمولة المطعم", "10٪ من قيمة الطعام بعد الخصم، في الطلبات المسلَّمة"],
            ["حصة من رسم التوصيل", "المنصة تحتفظ بـ 15٪ من رسم التوصيل؛ الكابتن 85٪"],
            ["إعلان الصدارة", "20 ₪ لليوم، باختيار المطعم ومراجعة الحوالة"],
            ["عضويات الزبائن", "50 ₪ أو 100 ₪ شهرياً حسب الباقة الفعلية"],
            ["كوبونات وتسعير نقاط", "أدوات تسويق مضبوطة من الإدارة دون كسر الهامش"],
        ],
    )
    para(
        doc,
        "للإدارة صفحة مالية تجمع الدخل والمصروف، عمولات المطاعم، تسوية ما يستحقه المطعم، وإعلانات الفترة. التسوية تُسجَّل حتى يبقى الحساب بين المنصة والمطبخ واضحاً.",
    )

    heading(doc, "17. ما الذي يظهر وما الذي يتوقف")
    bullets(
        doc,
        [
            "المطعم غير المعتمد: لا يظهر ولا يُطلب منه.",
            "المطعم المعتمد: ظاهر حتى يوقف الأدمن اللوحة.",
            "المطعم الموقوف: مخفي عن الزبائن؛ بياناته ومنيوه محفوظان.",
            "الإعلان المدفوع: ينتهي بانتهاء أيامه، دون أن يُخرج المطعم من المنصة.",
            "عضوية الزبون: تنتهي بانتهاء الشهر المدفوع ما لم تُجدَّد.",
            "كود الإحالة: لا ينتهي؛ البرنامج دائم.",
            "الطلب الجماعي: ينتهي إذا اكتمل أو أُلغي أو انتهت مهلة التجميع.",
        ],
    )

    heading(doc, "18. المبادئ التي لا نتنازل عنها")
    bullets(
        doc,
        [
            "غزة أولاً: مناطق، دفع، ولهجة المنتج لجمهور محلي.",
            "الزبون يشتري؛ الشريك والإدارة يديرون. لا خلط أدوار على نفس الحساب.",
            "لا إخفاء مطعم بسبب انتهاء عدّاد أيام بعد الاعتماد.",
            "العمولة والتوصيل والإعلان أرقام معلنة وقابلة للضبط من الإدارة دون غموض على الشريك.",
            "كل قرار إداري مهم له اسم وتاريخ في السجل.",
            "النمو الأصدق من صديق يدعو صديقاً، ويستمر طول عمر المنصة.",
            "الطعام الجماعي جزء من الثقافة، لذلك الطلب الجماعي خدمة أساسية لا حيلة تجميل.",
        ],
    )

    heading(doc, "19. الخلاصة")
    para(
        doc,
        "سفرة غزة هي السفرة الرقمية لأهل البلد: تكتشف المطعم، تطلب لوحدك أو مع جماعة، تدفع كما تعتاد، وتأخذ نقاطاً ترجع لك وجبة. المطعم يدخل بثقة الإدارة ويبقى، ويكسب من الطلبات ومن إعلان يشتريه إن أراد الصدارة. الكابتن يعمل بمسار واضح ومستحق معروف. والإدارة تشغّل اليوم كامل من غرفة واحدة، بفريق صلاحياته محددة وبأثر لكل تعديل.",
    )
    para(
        doc,
        "الفكرة ليست تطبيقاً يُنسخ من مدينة أخرى. الفكرة منصة ضيافة تليق بغزة: سريعة على الموبايل، محترمة للمطبخ المحلي، وواضحة في المال والعلاقة بين الأطراف الأربعة.",
        bold=True,
    )

    p = doc.add_paragraph()
    set_para_rtl(p, "center", space_before=28, space_after=4)
    add_text(p, "نهاية الوثيقة", size=12, color=MUTED)
    p2 = doc.add_paragraph()
    set_para_rtl(p2, "center")
    add_text(p2, "سفرة غزة — من غزة لأهلها", size=14, bold=True, color=PRIMARY)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    doc.save(str(OUT_AR))
    print(OUT)
    print(OUT_AR)


if __name__ == "__main__":
    build()
